from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "docs" / "reports"
RED = RGBColor(0xC0, 0x00, 0x00)


def all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def replace_text(paragraph, old, new):
    if old not in paragraph.text:
        return False
    full_text = paragraph.text.replace(old, new)
    paragraph.clear()
    run = paragraph.add_run(full_text)
    run.font.color.rgb = RED
    return True


def refresh_verification_metadata(doc, language):
    replacements = [
        ("https://github.com/SerikaKawano/UoL_CSM500-2026-APR", "https://github.com/SerikaKawano/PawMatch"),
    ]
    if language == "ja":
        replacements.extend(
            [
                ("22件成功、失敗0件", "39件成功、失敗0件"),
                ("22件の自動テスト", "39件の自動テスト"),
                ("2026年9月25日", "2026年9月28日"),
            ]
        )
    else:
        replacements.extend(
            [
                ("22 passed; 0 failed", "39 passed; 0 failed"),
                ("22 passing automated tests", "39 passing automated tests"),
                ("25 September 2026", "28 September 2026"),
            ]
        )
    changed = False
    for paragraph in all_paragraphs(doc):
        for old, new in replacements:
            changed = replace_text(paragraph, old, new) or changed
    return changed


def insert_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    result = paragraph._parent.add_paragraph()
    result._p.getparent().remove(result._p)
    new_p.getparent().replace(new_p, result._p)
    if style:
        result.style = style
    if text:
        run = result.add_run(text)
        run.font.color.rgb = RED
    return result


def replace_with_red(paragraph, text):
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.color.rgb = RED


def add_figure(after_paragraph, image_path, caption):
    image_p = insert_after(after_paragraph)
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.keep_with_next = True
    image_p.add_run().add_picture(str(image_path), width=Inches(6.45))
    caption_p = insert_after(image_p, caption, "Caption")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.keep_together = True
    return caption_p


def edit_report(path, language):
    doc = Document(str(path))
    metadata_changed = refresh_verification_metadata(doc, language)
    texts = [p.text for p in doc.paragraphs]
    marker = "本システムの基本的な運用単位" if language == "ja" else "The basic operating unit of the system"
    if any(marker in text for text in texts):
        if metadata_changed:
            doc.save(str(path))
        return

    abstract_heading = "Abstract" if language == "en" else "要旨"
    abstract_index = next(i for i, p in enumerate(doc.paragraphs) if p.text == abstract_heading)
    abstract_p = doc.paragraphs[abstract_index + 1]
    abstract_addition = (
        " The operating model preserves direct communication between the adopter and the rehomer, while PawMatch reviewers and administrators manage the more burdensome screening, evidence and workflow tasks."
        if language == "en"
        else " 運用上は里親希望者と譲渡者の直接対話を残しながら、負担の大きい審査、証拠確認、工程管理をPawMatchの審査担当者と管理者が担う。"
    )
    run = abstract_p.add_run(abstract_addition)
    run.font.color.rgb = RED

    anchor_text = (
        "The proposed system is a risk-aware decision-support prototype"
        if language == "en"
        else "提案するシステムは、犬猫の個人間譲渡における応募者審査を支える"
    )
    anchor = next(p for p in doc.paragraphs if p.text.startswith(anchor_text))
    if language == "ja":
        p1 = insert_after(anchor, "本システムの基本的な運用単位は、①里親希望者、②譲渡者、③審査担当者、④管理者の四つのロールである。ただし③と④はサービス提供側のPawMatch運営としてまとめて扱う。里親希望者は相談と必要情報の提出を行い、譲渡者は実際にペットを飼育している者として、性格、健康状態、日常のケア、譲渡条件などの質問に回答する。PawMatch運営は、適合性確認、書類、本人確認、審査工程、判断理由と監査記録を整理・管理する。", "Normal")
        p2 = insert_after(p1, "相談履歴は三者が参加する共有スレッドとする。ペット固有の事実は譲渡者へ直接確認できる一方、審査へ進める判断、追加書類の依頼、本人確認、工程の更新はPawMatch運営が担当する。管理者は通常の相談対応者ではないが、システム管理者として権限、設定、アカウント、監査を管理し、必要な場合は運営として記録を確認できる。この分担により、譲渡者の知識と意思決定への参加を失わずに、負担の大きい審査・手続き・確認をサービス側が支援する。", "Normal")
        caption = "図2a PawMatchの役割分担：里親希望者と譲渡者の直接対話を残し、審査・手続き・確認をPawMatch運営が支援する。出典：筆者作成。"
        image = REPORTS / "figures" / "role-model-ja.png"
    else:
        p1 = insert_after(anchor, "The basic operating unit of the system consists of four roles: (1) adopter, (2) rehomer, (3) reviewer and (4) administrator. Roles 3 and 4 are grouped as the service-side PawMatch operations team. The adopter asks questions and submits required information. The rehomer, who actually cares for the animal, answers questions about behaviour, health, daily care and adoption conditions. PawMatch operations organises and manages suitability review, documents, identity checks, workflow stages, decision rationale and audit records.", "Normal")
        p2 = insert_after(p1, "The enquiry history is a shared thread in which all three parties can participate. Pet-specific facts can be asked directly of the rehomer, whereas movement into formal screening, requests for additional evidence, identity checks and workflow updates are handled by PawMatch operations. The administrator is not normally the primary correspondent, but manages permissions, settings, accounts and audit access and may inspect records as part of operations. This division preserves the rehomer's knowledge and participation in the decision while transferring burdensome screening, procedure and verification work to the service provider.", "Normal")
        caption = "Figure 2a PawMatch responsibility model: direct adopter-rehomer dialogue is retained while PawMatch operations supports screening, procedures and verification. Source: author."
        image = REPORTS / "figures" / "role-model-en.png"
    add_figure(p2, image, caption)

    if language == "ja":
        decision_p = next(p for p in doc.paragraphs if p.text.startswith("判断支援システムは、情報を整理して人の判断を助け"))
        replace_with_red(decision_p, "判断支援システムは、情報を整理して人の判断を助け、自動化しないという本試作の目的に合う。PawMatch運営は審査の根拠と次の対応を整理・記録するが、最終的な譲渡には里親希望者と譲渡者の合意が必要である。譲渡判断は動物福祉、応募者の責任、個人情報に影響するため、この文脈で完全自動の承認・拒否は適切ではない。")
        implementation_p = next(p for p in doc.paragraphs if p.text.startswith("以下にコードの所在と実装した技術構成を示す"))
        insert_after(implementation_p, "相談スレッドでは、相談した里親希望者、対象ペットの譲渡者、PawMatch運営だけがメッセージを閲覧・送信できる。ペット固有の回答と運営上の案内は、発信者のロールが分かる表示で区別する。対象外の譲渡者は他者の相談にアクセスできず、審査工程の更新権限は従来どおり審査担当者と管理者に限定する。", "Normal")
    else:
        decision_p = next(p for p in doc.paragraphs if p.text.startswith("Decision Support Systems are relevant because"))
        replace_with_red(decision_p, "Decision Support Systems are relevant because the prototype organises information and assists judgement rather than automating it. PawMatch operations organises and records the evidence and next action, but final adoption requires agreement between the adopter and rehomer. Because adoption decisions affect animal welfare, applicant responsibility and personal information, fully automatic approval or rejection would be inappropriate in this context.")
        implementation_p = next(p for p in doc.paragraphs if p.text.startswith("The source repository and the implemented technology are shown below"))
        insert_after(implementation_p, "In the enquiry thread, only the adopter who opened the enquiry, the rehomer responsible for the listed pet and PawMatch operations can read and send messages. Role labels distinguish pet-specific answers from operational guidance. Unrelated rehomers cannot access another person's enquiry, while authority to advance formal screening stages remains restricted to reviewers and administrators.", "Normal")

    doc.save(str(path))


edit_report(REPORTS / "Final_Report_Integrated_Draft_JA.docx", "ja")
edit_report(REPORTS / "Final_Report_Integrated_Draft_EN.docx", "en")
