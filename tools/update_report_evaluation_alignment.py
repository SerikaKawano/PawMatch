from pathlib import Path

from docx import Document
from docx.shared import RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "docs" / "reports"
RED = RGBColor(255, 0, 0)


JA = {
    6: "本研究は、日本の犬猫譲渡におけるリスクを考慮した審査画面を開発・評価する。試作は動物固有の世話の条件と応募者の計画、確認状態、人による段階的判断を結びつける。申込情報のみの画面とPawMatchの確認支援付き画面を比べるシナリオ評価を設計し、重要リスクの発見、判断理由の質、情報の分かりやすさ、有用性、点数への過度な依存を調べる。画面操作の経過時間は、実際の譲渡で生じる追加連絡や書類確認を再現しないため、主成果ではなく補助的な観察値として扱う。現時点のソフトウェアと合成ケースは実装可能性を示すが、参加者評価と実際の譲渡成果は未収集である。本報告書は元の問題設定・設計・評価計画を維持しつつ、実装済みの機能と検証待ちの仮説を区別する。運用上は里親希望者と譲渡者の直接対話を残しながら、負担の大きい審査、証拠確認、工程管理をPawMatchの審査担当者と管理者が担う。",
    31: "先行研究やサービスは、譲渡までの道筋を伝えたり、団体の案件を管理したりする助けになる。本研究ではさらに、動物の必要条件、応募者の計画、未確認事項を一画面で示すと、審査担当者は重要なリスクを見つけ、次の対応を理由とともに説明しやすくなるかを確かめる。既存製品との優劣や、実際の譲渡件数の増加までは結論にしない。",
    33: "研究課題（RQ）は、申込情報のみの画面とPawMatchの確認支援付き画面を同じケースで使ったとき、参加者の判断にどのような違いが表れるかを、次の三点に整理したものである。",
    34: "・RQ1（重要な確認に気づけるか）：動物の必要なケアと応募者の計画を結び付け、事前に定めた重要リスク、未確認事項、次に確かめる内容を見落とさず挙げられるか。",
    35: "・RQ2（判断を説明しやすいか）：なぜその候補を選び、何を追加確認するのかを具体的に説明できるか。自由記述の研究者採点に加え、各課題後の有用性、情報の分かりやすさ、判断理由を説明できる自信を5段階で尋ねる。",
    36: "・RQ3（画面を過信しないか）：点数が高くても住居やケアの証拠が未確認のとき、数値だけで先へ進めず、確認待ちに気づけるか。重要リスクの見落としと、画面の結論・数値を重視した程度の自己評価を併せて確認する。",
    49: "評価のため、試作には二つの審査画面を用意する。「申込情報のみ」は同じ動物・応募者の事実と確認状態を示す。「PawMatch確認支援付き」は、それらに適合点、点数の根拠、リスク、次の確認を加える。どちらも候補順と元の情報を同じにし、確認支援を加えたときの判断と説明の違いを比較する。",
    68: "評価には合成の動物・応募者プロファイルを使うシナリオ課題を採用する。目的は技術的な動作確認だけではなく、試作が判断に必要な情報の理解と説明を支えるかを調べることである。参加者は「申込情報のみ」と「PawMatch確認支援付き」で同じ譲渡場面を検討する。",
    69: "「申込情報のみ」は応募者の事実と確認状態を単純なプロファイルとして示す。「PawMatch確認支援付き」には重み付き適合点、その根拠、リスク、推奨される次の確認を加える。参加者には各候補への対応、優先候補、関連するリスク、判断理由、次に確かめる内容を回答してもらう。",
    71: "主に測定するのは、想定した重要リスクの発見率、判断理由の明確さ、妥当な次の確認、有用性、情報の分かりやすさ、判断理由を説明できる自信、画面の結論や点数への過度な依存である。各課題後の5段階評価と自由記述を、同じ参加者・同じケースの組で比較する。画面操作の経過時間も記録するが、実際の譲渡で必要な追加連絡や書類確認を含まないため、補助的な観察値とする。",
    72: "図4に評価設計を示す。同じ合成の譲渡場面を「申込情報のみ」と「PawMatch確認支援付き」で検討し、重要リスクの発見、判断理由、次の確認、主観評価を比較する。表示順はAB／BAで交互にし、経過時間は参考値として記録する。",
    73: "図4 比較設計図予定：AB／BA順序、同じケースの反復による学習、重要リスクの発見、判断理由、課題後の5段階評価。Proposalの元図は付録Bに保持する。［判読可能な更新図を挿入］",
    75: "本評価は、同意を得た成人5名を目標とする形成的なテストである。Office for Health Improvement and Disparities（2020）は、課題を通じて使いにくさを見つける質的ユーザビリティテストには5～6名を勧める。一方、GOV.UK User Research Community（2018）は、サービス全体のユーザビリティを定量的にベンチマークする際には実際の利用者または利用見込み者30～60名の募集を勧める。5名は個人研究としての現実的な目標であり、統計的検出力に基づく人数でもサービス全体のベンチマークでもない。重要リスク、記述理由、課題後評価を中心に参加者ごとの観察を示し、統計的有意差、母集団での性能、実務上の審査時間短縮は主張しない。",
    76: "各参加者は合成の2ケースを両画面で扱い、計4課題に回答する。ABとBAの表示順を交互に割り当てるが、5名では3対2となり、同じケースの反復による学習も残る。参加者ごとに重要リスクの発見、具体的な操作上の問題、妥当な次の確認、記述理由、有用性、分かりやすさ、説明の自信、画面への依存を比較する。自由記述は事前の基準で判定し、分母、欠測、表示順、関連経験を報告する。経過時間は参考値として中断の有無とともに示す。この評価は設計上の問題と有望な兆候を探すもので、平均的な効果を推定しない。動物保護の実務者が参加しない場合、保護団体での使用を検証したとはいえない。",
    77: "［要検証：説明文を用意し、参加前に口頭またはEメールで同意を得た記録、募集方法、参加条件、人数、実施日、必要な倫理審査、固定したリスク基準と採点重みを記録する。アプリ上でも同意確認を行う。合成回答や架空の所要時間で代用しない。］",
    83: "最終提出では、最終報告書、GitHubの成果物、3分間の動画を、一つか二つの代表的なケースで構成する。実演では、申込情報だけでは整理しにくいリスクが、点数の根拠、警告、確認状態、推奨される対応によってどのように確認しやすくなるかを示す。",
    88: "外部の成人テスターには説明文を渡し、参加前に口頭またはEメールで同意を得る。アプリ上でも研究条件を確認し、同意したことを記録する。研究記録は参加者コードで扱い、撤回を可能にし、アクセス、保存期間、削除を文書化する。現行のデモログインは本番水準の本人確認ではなく、本研究では実物の本人確認書類、顔画像、収入証明を収集しない。",
    95: "表4は2026年9月28日にローカルの作業ツリーで行った検証を示す。42件の自動テストは、権限境界、相談情報の所有者、合成ケースの整合性、採点と確認状態の区別、管理者の重み変更の反映、研究参加の同意と撤回、人による最終審査を扱う。さらにBellaの1ケースで、里親希望者の相談、譲渡者とPawMatch運営の返信、適合性・書類確認、双方が確認する面談・トライアル・最終判断、譲渡承認までを一続きで検証し、完了工程に日時があり時系列が矛盾しないことを確認した。これはソフトウェア動作の確認であり、使いやすさ、公平性、動物福祉、実際の譲渡成果の証拠ではない。",
    99: "［結果待ち：依頼、同意、開始、完了、撤回の人数、重要リスクの発見率、妥当な次の確認、理由の判定、有用性、情報の分かりやすさ、説明の自信、画面への依存、欠測、対応のあるケースごとの比較を記入する。経過時間は中断を除外した補助値として示す。発言の引用には同意を得る。現時点で参加者の成果は主張しない。］",
    101: "図7 結果図予定：同じ参加者・同じケースについて、重要リスク発見率、有用性、情報の分かりやすさ、説明の自信、画面への依存の差を、分母と欠測とともに示す。経過時間は補助図または注記にする。［実際の参加者データ収集後にのみ挿入］",
    103: "想定する方法上の貢献は、同じ合成ケースについて二つの審査画面を比較可能にすることである。自己選択した5名では、研究者が作ったケースと採点基準、反復による学習、専門経験の違いが一人ひとりの結果に強く影響し得る。全員に同じ傾向が出ても、その5名と課題で観察されたことであり、保護団体全体への正確な効果推定ではない。リスク発見や説明が改善しても、自己評価だけで実務上の有効性が確定するわけではない。点数が不適切な前進を促したり、画面への依存を強めたりすれば設計を改める。経過時間は実際の追加連絡を含まないため効率の主証拠にしない。実際の譲渡成立、返還、動物福祉の結果には現場での追跡が必要であり、短期評価の範囲外である。",
    106: "合成データとシナリオを使うことは、倫理的リスクを抑えた評価設計にもつながる。申込情報のみの画面と比べ、重要リスクの発見、判断理由、次の確認、情報の分かりやすさ、認識された有用性のいずれかに一貫した改善があり、画面への依存や重大な見落としが増えなければ、この種の判断支援が安全で透明な譲渡の業務過程に役立つ可能性を示す材料になる。ただし実際の譲渡成果を直接示すものではない。",
    109: "本研究の中心仮説は、譲渡審査の簡略化とは確認を省くことではなく、動物ごとの世話の条件、応募者の計画、未確認の証拠、次に人が行う対応を一件の判断につなげることだ、というものである。試作はこの流れを実装し、42件の自動テストで一部の機能的性質を確認した。4ロールが参加する一匹の相談から譲渡承認までの合成シナリオも、権限と日時の整合性を保って完了できた。ただし、これだけで人の判断が改善したとはいえない。",
    110: "検証する予測は、申込情報のみの画面に比べ、PawMatchの確認支援付き画面では、参加者が事前に定めた重要リスクと次の確認を少なくとも同程度に発見し、判断理由をより具体的に記し、有用性、分かりやすさ、説明の自信を高く評価する、というものである。同時に、高得点でも住居やケアの証拠が未確認の事例で、重大な見落としや画面への依存が増えないことを確かめる。改善が一部の参加者やケースだけなら、その条件と反対例を結果として示す。経過時間は実務上の審査時間短縮を示さない補助値とする。［参加者評価後：差、分母、欠測、反対例を含め、仮説のどこが支持されたかを示す。］",
}


EN = {
    6: "This study develops and evaluates a risk-aware review interface for dog and cat rehoming in Japan. The prototype links animal-specific care requirements to applicant plans, verification states and a staged human decision. A scenario comparison between Application information only and With PawMatch review support examines recognition of important risks, quality of written reasons, information clarity, perceived usefulness and possible over-reliance on scores. Elapsed screen time is retained only as a secondary contextual observation because a prototype task cannot reproduce follow-up communication and document checks in a real adoption review. The current software and synthetic cases establish feasibility; participant results and real-world adoption outcomes have not yet been obtained. The report preserves the original problem, design and evaluation proposal while distinguishing implemented functions from hypotheses awaiting evaluation. The operating model preserves direct communication between the adopter and the rehomer, while PawMatch reviewers and administrators manage the more burdensome screening, evidence and workflow tasks.",
    31: "Existing work helps adopters understand the journey and organisations manage cases. This study asks a narrower unanswered question: when animal needs, applicant plans and pending checks appear together, can a reviewer identify important risks and explain the next action without relying only on the score? The study does not claim superiority over existing products or an increase in real adoptions.",
    33: "The research questions (RQs) organise three observable differences when the same case is reviewed with Application information only and With PawMatch review support:",
    34: "• RQ1 — Noticing important checks: Can participants connect the animal's care needs with the applicant's plan and identify the predefined important risks, missing evidence and appropriate next checks?",
    35: "• RQ2 — Explaining the decision: Can participants explain why they selected a candidate and what must be checked next? Researcher coding of written answers is combined with task-level 1–5 ratings of usefulness, information clarity and confidence in explaining the decision.",
    36: "• RQ3 — Avoiding over-reliance: When a provisional score is high but housing or care evidence is still missing, do participants notice the pending check rather than advancing solely because of the number? Critical misses are considered together with self-reported reliance on screen conclusions and scores.",
    49: "The evaluation uses two review screens. Application information only presents the same animal and applicant facts and verification states. With PawMatch review support adds the suitability score, reasons for the score, risks and suggested next checks. Candidate order and source facts remain identical, allowing the study to compare decisions and explanations when review support is added.",
    68: "The evaluation uses scenario-based tasks with synthetic animal and applicant profiles. Its purpose is not only to demonstrate technical operation, but to examine whether the prototype helps participants understand and explain information needed for review. Participants consider the same rehoming cases using Application information only and With PawMatch review support.",
    69: "Application information only presents applicant facts and verification states in a simple profile. With PawMatch review support adds a weighted suitability score, its explanation, risks and suggested next checks. For each candidate, participants record an action, select a priority candidate, identify risks, explain their decision and propose next checks.",
    71: "Primary measures are recall of predefined important risks, clarity of the written decision, appropriate next checks, perceived usefulness, information clarity, confidence in explaining the decision and possible over-reliance on screen conclusions or scores. Task-level 1–5 ratings and written answers are compared within the same participant and case. Elapsed screen time is also recorded, but is secondary because it excludes the follow-up communication and document checks required in real rehoming.",
    72: "Figure 4 shows the evaluation design. Participants consider the same synthetic case using Application information only and With PawMatch review support. Important-risk recall, written reasons, next checks and task-level ratings are compared. Display order alternates AB/BA; elapsed time is retained as contextual information.",
    73: "Figure 4 Planned comparison design: AB/BA order, repeated-case learning caveat, important-risk recall, written reasons and task-level ratings. The original Proposal comparison figure is preserved in Appendix B. [INSERT UPDATED LEGIBLE DIAGRAM]",
    75: "This formative evaluation targets five consenting adult participants. The Office for Health Improvement and Disparities (2020) recommends five to six people for qualitative usability testing intended to identify problems in tasks. By contrast, the GOV.UK User Research Community (2018) advises recruiting 30 to 60 actual or likely users for service-wide usability benchmarking. Five participants are a pragmatic choice for this individual project, not a powered sample or a service benchmark. Participant-level observations of important risks, written reasons and task ratings will be primary. The study will not claim statistical significance, population performance or a real-world reduction in review time.",
    76: "Each participant will see two selected synthetic cases in both modes, making four tasks. AB and BA display orders will be alternated; with five people the split is necessarily three versus two, and repeating cases may teach participants what to look for. Within each participant, the study compares important-risk recall, concrete usability problems, appropriate next checks, written reasons, usefulness, clarity, confidence and reliance on the screen. Free text will be coded against a predefined rubric, with denominators, missing data, order and relevant experience disclosed. Elapsed time will be reported only as contextual information alongside interruptions. The study will look for design problems and promising signals, not estimate an average treatment effect. If no animal-welfare practitioner participates, the findings cannot validate use by rehoming organisations.",
    77: "[PENDING VALIDATION: prepare participant information and record oral or email consent before participation; also retain the in-application consent confirmation. Document recruitment, eligibility, participant count, dates, ethics approval if required, the frozen risk rubric and scoring weights. Do not substitute simulated participants or invented times.]",
    83: "For the final submission, the report, GitHub artefacts and three-minute video will be structured around one or two representative scenarios. The demonstration should show how risks that are difficult to organise from application information alone become easier to inspect through score explanations, warnings, verification states and suggested actions.",
    88: "External adult testers must receive clear participant information and provide oral or email consent before participating. The application also records their confirmation of the study conditions. Research records use participant codes, support withdrawal and apply documented access, retention and deletion rules. The demo login is not production-grade identity verification; this study does not collect real identity documents, facial images or income evidence.",
    95: "Table 4 reports checks run on the local working tree on 28 September 2026. The 42 automated tests cover role boundaries, consultation ownership, synthetic-case consistency, separation of scores from verification, propagation of administrator weights, consent and withdrawal in study sessions, and human-gated final review. One complete Bella scenario also checks the adopter's enquiry, replies by the rehomer and PawMatch operations, suitability and document review, joint meeting, trial and final sign-off, adoption approval, and chronological dates for every completed stage. These checks establish selected software behaviour; they do not establish usability, fairness, animal-welfare impact or adoption outcomes.",
    99: "[PENDING RESULTS: report invited, consented, started, completed and withdrawn counts; important-risk recall; appropriate next checks; coded explanations; usefulness; information clarity; confidence in explaining decisions; screen reliance; missing data; and paired case-level comparisons. Report elapsed time only as a secondary value excluding interrupted tasks. Quote participants only with consent. No participant outcome is claimed at this stage.]",
    101: "Figure 7 Planned results figure: within-participant, same-case differences in important-risk recall, usefulness, information clarity, confidence and screen reliance, with denominators and missing data shown. Elapsed time will appear only in a secondary figure or note. [INSERT ONLY AFTER REAL PARTICIPANT DATA ARE COLLECTED]",
    103: "The anticipated methodological contribution is an inspectable comparison of two review screens on the same synthetic cases. With five self-selected participants, researcher-written cases and scoring rules, learning between repeated tasks and differences in domain experience can strongly influence every observation. Even a unanimous pattern would be a finding about these tasks and people, not a precise estimate for rehoming organisations. Improved risk recall or explanations would not by itself establish operational effectiveness. If scores encourage unjustified advancement or greater reliance on the screen, the design should be revised. Elapsed time excludes real follow-up and is not primary evidence of efficiency. Real placement success, return and animal welfare require field follow-up and are outside this short evaluation.",
    106: "Synthetic data and scenarios also provide an ethically lower-risk evaluation design. If, compared with Application information only, review support shows a consistent improvement in important-risk recall, written reasons, next checks, information clarity or perceived usefulness without increasing critical misses or screen reliance, the findings can provide evidence of potential value for safer and more transparent rehoming workflows. They will not directly establish real adoption outcomes.",
    109: "This study advances a specific design hypothesis: rehoming review may be simplified not by omitting checks, but by bringing each animal's care needs, the applicant's plan, outstanding evidence and the next human action into one traceable decision. The prototype implements that workflow, and 42 automated tests establish selected functional properties. A synthetic journey involving all four roles also completes one pet's enquiry through adoption approval while preserving permissions and chronological dates. These results do not show that people make better decisions with it.",
    110: "The testable prediction is that, compared with Application information only, With PawMatch review support will help participants identify at least as many predefined important risks and next checks, provide more specific reasons, and rate usefulness, clarity and confidence more highly. At the same time, high-scoring cases with unresolved housing or care evidence test whether critical misses or screen reliance increase. If improvement appears only for some people or cases, the report will state those conditions and counterexamples. Elapsed time is a secondary value and will not be interpreted as real-world review-time reduction. [After participant testing: insert differences, denominators, missing data and counterexamples, and state which parts of the prediction were supported.]",
}


def replace_paragraph(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.color.rgb = RED


def replace_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    replace_paragraph(paragraph, text)
    for extra in cell.paragraphs[1:]:
        extra.clear()


def update(path: Path, replacements: dict[int, str], language: str) -> None:
    document = Document(path)
    for index, text in replacements.items():
        replace_paragraph(document.paragraphs[index], text)

    checks = document.tables[3]
    replace_cell(checks.cell(1, 1), "42件成功、失敗0件" if language == "ja" else "42 passed; 0 failed")
    replace_cell(
        checks.cell(1, 2),
        "役割、所有権、研究セッション、採点、審査フロー、全ロール横断の譲渡完了。"
        if language == "ja"
        else "Roles, ownership, research sessions, scoring, review flow and an all-role completed adoption journey.",
    )

    outcomes = document.tables[4]
    labels_ja = [
        ("評価項目", "申込情報のみ", "PawMatch確認支援付き"),
        ("完了した課題", "［件数／予定件数］", "［件数／予定件数］"),
        ("重要リスク発見率", "［発見数／機会数］", "［発見数／機会数］"),
        ("判断理由の明確さ", "［平均0～3・採点数］", "［平均0～3・採点数］"),
        ("有用性・分かりやすさ・説明の自信", "［各平均1～5］", "［各平均1～5］"),
        ("画面の結論・数値への依存", "［平均1～5］", "［平均1～5］"),
    ]
    labels_en = [
        ("Outcome", "Application information only", "With PawMatch review support"),
        ("Completed tasks", "[n / planned tasks]", "[n / planned tasks]"),
        ("Important-risk recall", "[detected / opportunities]", "[detected / opportunities]"),
        ("Clarity of written reasons", "[mean 0–3; coded n]", "[mean 0–3; coded n]"),
        ("Usefulness, clarity and confidence", "[each mean 1–5]", "[each mean 1–5]"),
        ("Reliance on screen conclusions or scores", "[mean 1–5]", "[mean 1–5]"),
    ]
    for row, values in zip(outcomes.rows, labels_ja if language == "ja" else labels_en):
        for cell, value in zip(row.cells, values):
            replace_cell(cell, value)

    document.save(path)


if __name__ == "__main__":
    update(REPORTS / "Final_Report_Integrated_Draft_JA.docx", JA, "ja")
    update(REPORTS / "Final_Report_Integrated_Draft_EN.docx", EN, "en")
    print("Updated Japanese and English reports.")
